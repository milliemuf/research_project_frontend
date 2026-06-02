"""Regenerate SUPPLEMENTARY_reproduction.docx from SUPPLEMENTARY_reproduction.md.
Handles headings, paragraphs, **bold**, bullet lists, markdown tables, and fenced
code blocks (simple ``` toggle; monospace)."""
from __future__ import annotations
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent
SRC = ROOT / "SUPPLEMENTARY_reproduction.md"
DST = ROOT / "SUPPLEMENTARY_reproduction.docx"

lines = SRC.read_text(encoding="utf-8").split("\n")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(11)
for sec in doc.sections:
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.2)

_inline = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")

def add_runs(p, txt):
    pos = 0
    for m in _inline.finditer(txt):
        if m.start() > pos:
            p.add_run(txt[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            p.add_run(tok[2:-2]).bold = True
        else:
            r = p.add_run(tok[1:-1]); r.font.name = "Consolas"
        pos = m.end()
    if pos < len(txt):
        p.add_run(txt[pos:])

def is_table_sep(s):
    return bool(re.match(r"^\s*\|?[\s:\-|]+\|[\s:\-|]*$", s)) and "-" in s

i = 0
n = len(lines)
while i < n:
    line = lines[i].rstrip()
    s = line.strip()

    # fenced code block (toggle on any line starting with ```)
    if s.startswith("```"):
        i += 1
        buf = []
        while i < n and not lines[i].strip().startswith("```"):
            buf.append(lines[i]); i += 1
        i += 1  # consume closing fence
        p = doc.add_paragraph()
        run = p.add_run("\n".join(buf)); run.font.name = "Consolas"; run.font.size = Pt(9)
        continue

    if not s:
        i += 1; continue

    # markdown table: a | row followed by a |---| separator
    if s.startswith("|") and i + 1 < n and is_table_sep(lines[i + 1]):
        header = [c.strip() for c in s.strip("|").split("|")]
        i += 2  # skip header + separator
        rows = []
        while i < n and lines[i].strip().startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]; c.text = ""; c.paragraphs[0].add_run(h).bold = True
        for row in rows:
            cells = t.add_row().cells
            for j in range(min(len(row), len(header))):
                cells[j].text = row[j]
        continue

    if s.startswith("# "):
        doc.add_heading(s[2:].strip(), level=0); i += 1; continue
    if s.startswith("## "):
        doc.add_heading(s[3:].strip(), level=1); i += 1; continue
    if s.startswith("### "):
        doc.add_heading(s[4:].strip(), level=2); i += 1; continue
    if re.match(r"^[-*]\s+", s):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, re.sub(r"^[-*]\s+", "", s)); i += 1; continue

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, s); i += 1

doc.save(DST)
print("wrote", DST, "paras:", len(doc.paragraphs), "tables:", len(doc.tables))
