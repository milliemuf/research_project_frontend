"""Build the non-anonymised technical paper for the supervisor: take the revised
manuscript and restore the author block (name, affiliation, e-mails) into the
blank paragraphs left by anonymisation, then export to PDF. Front matter only --
no response letter or supplementary -- so the body word count is unchanged."""
import os
import docx
import docx2pdf
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(ROOT, "MS_R1_revised_clean.docx")
DOCX = os.path.join(ROOT, "Technical_Paper_for_Supervisor.docx")
PDF = os.path.join(ROOT, "Technical_Paper_for_Supervisor.pdf")

AUTHORS = "Millicent Mufambi*, W. Makondo"
AFFIL = ("Department of Software Engineering, Harare Institute of Technology, "
         "P. O. Box BE 277, Belvedere, Harare, Zimbabwe")
EMAILS = "h240624a@hit.ac.zw; wmakondo@hit.ac.zw"

d = docx.Document(SRC)
# P0 = title; P1..P5 = blank author area left by anonymisation. Fill the first three.
blanks = [p for p in d.paragraphs[1:6] if not p.text.strip()]
fills = [(AUTHORS, True), (AFFIL, False), (EMAILS, False)]
for para, (text, bold) in zip(blanks, fills):
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.bold = bold
d.save(DOCX)
print("wrote", DOCX)

docx2pdf.convert(DOCX, PDF)
print("wrote", PDF, "exists:", os.path.exists(PDF))
