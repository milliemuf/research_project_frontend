"""Remove em-dashes (and spaced en-dashes used as clause separators) from .docx
deliverables, replacing them with a comma so sentences still read cleanly. Hyphens
and unspaced en-dashes in ranges (e.g. 2.5-10%) are left alone."""
import re, sys
import docx

DOCS = [
    "Technical_Paper_for_Supervisor.docx",
    "MS_R1_revised_clean.docx",
    "MS_R1_baseline_anon.docx",
    "AUTHOR_RESPONSE_LETTER.docx",
    "SUPPLEMENTARY_reproduction.docx",
    "COMBINED_for_supervisor.docx",
    "MS_R1_tracked_changes.docx",
    "PROJECT MTECH Software Engineering - COMPLETED.docx",
]

def fix(t):
    if "—" not in t and "–" not in t:
        return t
    t = re.sub(r"\s*—\s*", ", ", t)        # em-dash -> comma
    t = re.sub(r"\s+–\s+", ", ", t)        # spaced en-dash -> comma (ranges keep hyphen)
    t = re.sub(r",\s*,", ",", t)                 # tidy any double commas
    t = re.sub(r"\s+,", ",", t)                  # tidy " ," -> ","
    return t

def strip_doc(path):
    d = docx.Document(path)
    n = 0
    def do_runs(runs):
        nonlocal n
        for r in runs:
            if "—" in r.text or "–" in r.text:
                new = fix(r.text)
                if new != r.text:
                    r.text = new; n += 1
    for p in d.paragraphs:
        do_runs(p.runs)
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    do_runs(p.runs)
    # headers/footers
    for sec in d.sections:
        for hf in (sec.header, sec.footer):
            for p in hf.paragraphs:
                do_runs(p.runs)
    d.save(path)
    return n

if __name__ == "__main__":
    targets = sys.argv[1:] or DOCS
    for f in targets:
        try:
            print(f"{strip_doc(f):4d} runs fixed in  {f}")
        except Exception as e:
            print(f"  skip {f}: {e}")
