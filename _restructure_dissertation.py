"""Restructure the dissertation into one cohesive document (no 'Appendix' sections):
  - delete the code listings (G), screenshots (E) and the redundant glossary (D);
  - move the algorithm spec (C), bug catalogues (A, B) into Chapter 3 as flowing
    sub-sections, and the per-case results (F) into Chapter 4;
  - references therefore become the last section;
  - add an explained PBFT-adaptation algorithm with equations into the methodology.
Operates in place on the COMPLETED .docx (original preserved separately)."""
import docx
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

DOC = "PROJECT MTECH Software Engineering - COMPLETED.docx"
d = docx.Document(DOC)


def find_h1(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix) and p.style.name == "Heading 1":
            return p
    return None


def find_h(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


def set_text(p, text):
    for ch in list(p._p):
        if ch.tag in (qn("w:r"), qn("w:hyperlink")):
            p._p.remove(ch)
    p.add_run(text)


def block_elements(hp):
    """Heading element + following siblings up to (not incl.) the next Heading 1."""
    els = [hp._p]
    el = hp._p.getnext()
    while el is not None:
        if el.tag == qn("w:p") and Paragraph(el, hp._parent).style.name == "Heading 1":
            break
        els.append(el)
        el = el.getnext()
    return els


def delete_block(prefix):
    hp = find_h1(prefix)
    if hp is None:
        print("  (delete) not found:", prefix); return
    for el in block_elements(hp):
        el.getparent().remove(el)
    print("  deleted:", prefix[:30])


def move_block(block_prefix, target_prefix, new_text, new_style="Heading 21"):
    hp = find_h1(block_prefix)
    tgt = find_h(target_prefix)
    if hp is None or tgt is None:
        print("  (move) not found:", block_prefix, "->", target_prefix); return
    els = block_elements(hp)
    set_text(hp, new_text)
    hp.style = d.styles[new_style]
    for el in els:
        tgt._p.addprevious(el)
    print("  moved:", block_prefix[:24], "->", new_text[:30])


# 1. delete code, screenshots, redundant glossary (Section 1.8 already defines terms)
delete_block("APPENDIX G")
delete_block("APPENDIX E")
delete_block("APPENDIX D")

# 2. move algorithm + catalogues into Chapter 3 (before Chapter 4)
move_block("APPENDIX C", "CHAPTER 4", "3.9 PBFT Adaptation: Algorithm and Equations")
move_block("APPENDIX A", "CHAPTER 4", "3.10 Synthetic Bug Catalogue (n = 40)")
move_block("APPENDIX B", "CHAPTER 4", "3.11 E-Commerce Bug Catalogue (n = 30)")

# 3. move per-case results into Chapter 4 (before 4.5 Proposed Model)
move_block("APPENDIX F", "4.5 Proposed Model", "4.4.10 Full Per-Case Experimental Results")

# 4. add an explained PBFT-adaptation algorithm with equations into 3.9
algo = find_h("3.9 PBFT Adaptation")
if algo is not None:
    paras = [
        ("The consensus layer adapts Practical Byzantine Fault Tolerance (PBFT) to a "
         "population of stochastic LLM agents. With a target fault tolerance of f, the "
         "replication bound fixes the validator count at n = 3f + 1 and the agreement "
         "quorum at 2f + 1; for the default f = 1 this gives n = 4 independent validators "
         "and a quorum of 3. The analyzer and the healer do not vote, so the voting "
         "population is genuinely independent."),
        ("For a candidate fix x the primary (healer) issues a pre-prepare carrying the byte "
         "digest D(x) of the proposal. Each validator v independently evaluates x and casts a "
         "prepare vote b_v in {accept, reject}. A replica becomes prepared on digest D when "
         "it holds the matching pre-prepare and at least 2f + 1 prepare messages agreeing on "
         "D, that is when |{v : prepare(v, D)}| >= 2f + 1. A prepared replica broadcasts a "
         "commit, and commits the fix once it has collected 2f + 1 matching commits: "
         "committed(D) iff |{v : commit(v, D)}| >= 2f + 1 and the replica is prepared on D. "
         "Because any two quorums of size 2f + 1 drawn from 3f + 1 replicas intersect in at "
         "least f + 1 replicas, and at least one of those is honest, no two honest replicas "
         "can commit different digests for the same slot; this is the safety property the "
         "protocol guarantees."),
        ("The procedure is summarised as follows. Phase 1 (pre-prepare): the primary "
         "broadcasts (D(x), x). Phase 2 (prepare): each validator verifies the signature and "
         "the digest, casts accept or reject, and broadcasts its vote. Phase 3 (commit): on "
         "reaching the 2f + 1 prepare quorum a validator broadcasts commit; on reaching the "
         "2f + 1 commit quorum it decides. An approved fix is then executed in the sandbox, "
         "which is the final gate. If fewer than 2f + 1 validators accept, the fix is "
         "rejected; if the quorum cannot form (for example under more than f faults or a "
         "network partition), the protocol makes no progress rather than committing an "
         "uncertified fix, preserving safety at the cost of liveness until recovery."),
    ]
    nxt = algo._p.getnext()
    anchor = Paragraph(nxt, algo._parent) if nxt is not None else None
    for text in reversed(paras):
        from docx.oxml import OxmlElement
        np = OxmlElement("w:p")
        algo._p.addnext(np)
        P = Paragraph(np, algo._parent)
        P.style = d.styles["Normal"]
        P.add_run(text)

d.save(DOC)
print("restructured ->", DOC)
# report heading-1 order
d2 = docx.Document(DOC)
print("=== Heading-1 order now ===")
for p in d2.paragraphs:
    if p.style.name == "Heading 1" and p.text.strip():
        print("  ", p.text.strip()[:50])
