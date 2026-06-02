"""Expand the dissertation chapters with substantive prose toward ~30,000 words.
Inserts Normal-style paragraphs at the END of named sub-sections (immediately
before the next heading), keeping the existing tables/figures in place. No
em-dashes are used in any inserted text. Run repeatedly with different SECTIONS
payloads (one chapter group per pass)."""
import sys, importlib
import docx
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DOC = "PROJECT MTECH Software Engineering - COMPLETED.docx"
HEAD = {"Heading 1", "Heading 21", "Heading 31"}


def _find_heading(d, prefix):
    for p in d.paragraphs:
        if p.style.name in HEAD and p.text.strip().startswith(prefix):
            return p
    return None


def _next_heading_el(p):
    el = p._p.getnext()
    while el is not None:
        if el.tag.endswith("}p"):
            para = Paragraph(el, p._parent)
            if para.style.name in HEAD and para.text.strip():
                return el
        el = el.getnext()
    return None


def insert_at_end_of(d, prefix, paras):
    h = _find_heading(d, prefix)
    if h is None:
        print("  NOT FOUND:", prefix); return 0
    anchor = _next_heading_el(h)
    n = 0
    for text in paras:
        np = OxmlElement("w:p")
        if anchor is not None:
            anchor.addprevious(np)
        else:
            d.element.body.append(np)
        P = Paragraph(np, h._parent)
        P.style = d.styles["Normal"]
        P.add_run(text)
        n += len(text.split())
    print(f"  +{n:5d}w  {prefix[:50]}")
    return n


def run(SECTIONS):
    d = docx.Document(DOC)
    total = 0
    for prefix, paras in SECTIONS:
        total += insert_at_end_of(d, prefix, paras)
    d.save(DOC)
    print(f"=== inserted {total} words ===")


if __name__ == "__main__":
    mod = importlib.import_module(sys.argv[1])
    run(mod.SECTIONS)
