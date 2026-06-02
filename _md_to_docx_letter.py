"""Regenerate AUTHOR_RESPONSE_LETTER.docx from AUTHOR_RESPONSE_LETTER.md."""
from __future__ import annotations
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent
SRC = ROOT / "AUTHOR_RESPONSE_LETTER.md"
DST = ROOT / "AUTHOR_RESPONSE_LETTER.docx"

text = re.sub(r"<!--.*?-->", "", SRC.read_text(encoding="utf-8"), flags=re.DOTALL)
lines = text.split("\n")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
for sec in doc.sections:
    sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)

_inline = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")

def add_runs(p, txt):
    """Render a markdown line into runs, honouring **bold**, *italic*, `code`."""
    pos = 0
    for m in _inline.finditer(txt):
        if m.start() > pos:
            p.add_run(txt[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            p.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*"):
            p.add_run(tok[1:-1]).italic = True
        else:  # `code`
            r = p.add_run(tok[1:-1]); r.font.name = "Consolas"
        pos = m.end()
    if pos < len(txt):
        p.add_run(txt[pos:])

para_buf: list[str] = []

def flush():
    global para_buf
    if not para_buf:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, " ".join(para_buf).strip())
    para_buf = []

for raw in lines:
    line = raw.rstrip()
    if not line.strip():
        flush(); continue
    if line.startswith("# "):
        flush()
        h = doc.add_heading(line[2:].strip(), level=0)
        continue
    if line.startswith("## "):
        flush()
        doc.add_heading(line[3:].strip(), level=1); continue
    if line.startswith("### "):
        flush()
        doc.add_heading(line[4:].strip(), level=2); continue
    if line.strip() == "---":
        flush(); continue  # section divider; blank gap is enough
    if re.match(r"^\s*-\s+", line):
        flush()
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, re.sub(r"^\s*-\s+", "", line)); continue
    para_buf.append(line.strip())

flush()
doc.save(DST)
print("wrote", DST)
